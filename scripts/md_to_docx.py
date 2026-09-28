"""Markdown → .docx для тех, кто правит текст в Word (защита, отчёты).

Поддерживает то, что реально встречается в наших документах: заголовки, списки, таблицы,
цитаты (реплики выступающего), горизонтальные линии, жирный шрифт, курсив и `код`.

    python scripts/md_to_docx.py docs/pitch_script.md outputs/Защита.docx
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INK = RGBColor(0x1B, 0x1D, 0x21)
MUTED = RGBColor(0x5A, 0x62, 0x70)
ACCENT = RGBColor(0x2F, 0x5B, 0xD0)
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_runs(par, text, size=11, color=INK):
    """Разбирает **жирный**, *курсив* и `код` внутри строки."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        run = par.add_run(chunk.strip("*`") if chunk[0] in "*`" else chunk)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if chunk.startswith("**"):
            run.bold = True
        elif chunk.startswith("`"):
            run.font.name = "Menlo"
            run.font.size = Pt(size - 1)
        elif chunk.startswith("*"):
            run.italic = True
    return par


def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def quote(doc, lines):
    """Реплика выступающего: отступ, курсив, полоска слева."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(24)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    add_runs(par, " ".join(lines), size=11.5, color=INK)
    for run in par.runs:
        run.italic = True
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8"); left.set(qn("w:color"), "2F5BD0")
    pbdr.append(left)
    par._p.get_or_add_pPr().append(pbdr)


def table(doc, rows):
    header, *body = rows
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, text in enumerate(header):
        cell = t.cell(0, j)
        cell.text = ""
        add_runs(cell.paragraphs[0], text, size=10)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade(cell, "EAF0FB")
    for i, row in enumerate(body, start=1):
        for j, text in enumerate(row[:len(header)]):
            cell = t.cell(i, j)
            cell.text = ""
            add_runs(cell.paragraphs[0], text, size=10)
    doc.add_paragraph()


def convert(md_path, out_path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    lines = Path(md_path).read_text().split("\n")
    i, pending_quote, pending_table = 0, [], []

    def flush():
        nonlocal pending_quote, pending_table
        if pending_quote:
            quote(doc, pending_quote); pending_quote = []
        if pending_table:
            table(doc, pending_table); pending_table = []

    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("> "):
            pending_quote.append(line[2:]); i += 1; continue
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):          # строка-разделитель
                pending_table.append(cells)
            i += 1
            continue
        flush()
        if not line.strip():
            i += 1; continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("# ").strip()
            h = doc.add_heading(level=min(level, 4))
            add_runs(h, text, size={1: 18, 2: 15, 3: 13, 4: 12}.get(level, 12),
                     color=INK if level > 1 else ACCENT)
            for run in h.runs:
                run.bold = True
        elif line.strip() in ("---", "***", "___"):
            par = doc.add_paragraph()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "D8DCE2")
            pbdr.append(bottom)
            par._p.get_or_add_pPr().append(pbdr)
        elif re.match(r"^\s*[-*]\s+", line):
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            par = doc.add_paragraph(style="List Number")
            add_runs(par, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            run = par.add_run("\n".join(block))
            run.font.name = "Menlo"; run.font.size = Pt(9.5); run.font.color.rgb = MUTED
        else:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(par, line)
        i += 1
    flush()

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source", type=Path)
    ap.add_argument("target", type=Path)
    args = ap.parse_args()
    out = convert(args.source, args.target)
    print(f"{args.source} → {out}")


if __name__ == "__main__":
    main()
